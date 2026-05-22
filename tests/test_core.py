from __future__ import annotations

import os
import json
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest import mock

from docx import Document
from fastapi import HTTPException

import app as app_module
from agent import chain, retriever, sessions
from knowledge import library


class IsolatedSessionsMixin:
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.data_dir = Path(self.tmp.name)
        self.original_session_paths = {
            "DATA_DIR": sessions.DATA_DIR,
            "SESSIONS_JSON_PATH": sessions.SESSIONS_JSON_PATH,
            "SESSIONS_DB_PATH": sessions.SESSIONS_DB_PATH,
            "SESSIONS_LEGACY_MIGRATED_PATH": sessions.SESSIONS_LEGACY_MIGRATED_PATH,
            "_INITIALIZED": sessions._INITIALIZED,
        }
        sessions.DATA_DIR = self.data_dir
        sessions.SESSIONS_JSON_PATH = self.data_dir / "sessions.json"
        sessions.SESSIONS_DB_PATH = self.data_dir / "sessions_store.sqlite3"
        sessions.SESSIONS_LEGACY_MIGRATED_PATH = self.data_dir / "sessions_legacy_imported.flag"
        sessions._INITIALIZED = False

    def tearDown(self) -> None:
        sessions.DATA_DIR = self.original_session_paths["DATA_DIR"]
        sessions.SESSIONS_JSON_PATH = self.original_session_paths["SESSIONS_JSON_PATH"]
        sessions.SESSIONS_DB_PATH = self.original_session_paths["SESSIONS_DB_PATH"]
        sessions.SESSIONS_LEGACY_MIGRATED_PATH = self.original_session_paths["SESSIONS_LEGACY_MIGRATED_PATH"]
        sessions._INITIALIZED = self.original_session_paths["_INITIALIZED"]
        self.tmp.cleanup()


class SessionStorageTests(IsolatedSessionsMixin, unittest.TestCase):
    def test_append_turn_persists_messages_and_summary(self) -> None:
        saved = sessions.append_turn(
            None,
            user_message="讲一下 BFS",
            assistant_message="BFS 适合按层遍历图。",
            assistant_sources=["graph_bfs_dfs.md"],
            assistant_evidence=[{"source": "graph_bfs_dfs.md", "excerpt": "BFS 使用队列。"}],
            assistant_used_rag=True,
        )

        loaded = sessions.get_session(str(saved["id"]))

        self.assertIsNotNone(loaded)
        assert loaded is not None
        self.assertEqual(len(loaded["messages"]), 2)
        self.assertEqual(loaded["messages"][1]["sources"], ["graph_bfs_dfs.md"])
        self.assertTrue(loaded["messages"][1]["used_rag"])
        self.assertIn("BFS", loaded["summary"])


class DocumentParsingTests(unittest.TestCase):
    def test_docx_upload_parsing_builds_preview_chunks(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            doc_path = Path(tmp) / "algorithm.docx"
            document = Document()
            document.add_paragraph("动态规划适合处理具有重叠子问题和最优子结构的问题。")
            document.add_paragraph("状态定义、状态转移和初始化是 DP 讲解的核心。")
            document.save(doc_path)

            parsed = library.parse_uploaded_document(doc_path.name, doc_path.read_bytes())

        self.assertEqual(parsed["file_name"], "algorithm.docx")
        self.assertGreater(parsed["chunk_count"], 0)
        self.assertGreater(parsed["text_quality"], 0)
        self.assertTrue(parsed["preview_chunks"])

    def test_markdown_upload_parsing_strips_basic_markup(self) -> None:
        content = b"# \xe5\x89\x8d\xe7\xbc\x80\xe5\x92\x8c\n\n- \xe7\x94\xa8\xe4\xba\x8e\xe5\xbf\xab\xe9\x80\x9f\xe6\xb1\x82\xe5\x8c\xba\xe9\x97\xb4\xe5\x92\x8c\n\n`sum[i] = sum[i-1] + a[i]`"

        parsed = library.parse_uploaded_document("prefix_sum.md", content)

        self.assertEqual(parsed["mime_type"], "text/markdown")
        self.assertIn("前缀和", parsed["extracted_text"])
        self.assertIn("用于快速求区间和", parsed["extracted_text"])
        self.assertGreater(parsed["chunk_count"], 0)

    def test_latex_upload_parsing_extracts_readable_text(self) -> None:
        content = r"""
        \section{动态规划}
        动态规划适合处理重叠子问题。
        \begin{itemize}
        \item 状态定义
        \item 状态转移
        \end{itemize}
        $dp[i] = \min(dp[i-1], dp[i-2])$
        """.encode("utf-8")

        parsed = library.parse_uploaded_document("dp.tex", content)

        self.assertEqual(parsed["mime_type"], "application/x-latex")
        self.assertIn("动态规划", parsed["extracted_text"])
        self.assertIn("状态定义", parsed["extracted_text"])
        self.assertNotIn("min(dp", parsed["extracted_text"])

    def test_pptx_upload_parsing_uses_python_pptx_when_available(self) -> None:
        fake_pptx = types.ModuleType("pptx")

        class FakeShape:
            def __init__(self, text: str) -> None:
                self.text = text

        class FakeSlide:
            def __init__(self, shapes) -> None:
                self.shapes = shapes

        class FakePresentation:
            def __init__(self, _stream) -> None:
                self.slides = [
                    FakeSlide([FakeShape("BFS 适合按层遍历。"), FakeShape("")]),
                    FakeSlide([FakeShape("动态规划关注状态转移。")]),
                ]

        fake_pptx.Presentation = FakePresentation

        with mock.patch.dict(sys.modules, {"pptx": fake_pptx}):
            parsed = library.parse_uploaded_document("algo.pptx", b"fake-pptx")

        self.assertEqual(
            parsed["mime_type"],
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
        self.assertIn("BFS", parsed["extracted_text"])
        self.assertIn("状态转移", parsed["extracted_text"])

    def test_ppt_upload_reports_conversion_hint(self) -> None:
        with self.assertRaisesRegex(ValueError, "另存为 \\.pptx"):
            library.parse_uploaded_document("legacy.ppt", b"legacy")

    def test_pdf_without_tesseract_falls_back_to_text_extraction_with_warning(self) -> None:
        blocks = [library.TextBlock(text="前缀和可以快速求区间和。", location="第 1 页")]

        with mock.patch("knowledge.library._extract_pdf_text_blocks", return_value=blocks), mock.patch(
            "knowledge.library._ocr_available",
            return_value=True,
        ), mock.patch(
            "knowledge.library._text_quality_profile",
            return_value=(0.2, 0.5),
        ), mock.patch(
            "knowledge.library._extract_pdf_blocks_via_ocr",
            side_effect=RuntimeError("已安装 OCR Python 依赖，但系统未安装 Tesseract 可执行程序。"),
        ):
            parsed = library.parse_uploaded_document("scan.pdf", b"fake-pdf")

        self.assertEqual(parsed["extraction_mode"], "text")
        self.assertFalse(parsed["ocr_used"])
        self.assertIn("Tesseract", parsed["extraction_warning"])
        self.assertIn("前缀和", parsed["extracted_text"])

    def test_pdf_without_tesseract_and_without_text_raises_clear_error(self) -> None:
        with mock.patch("knowledge.library._extract_pdf_text_blocks", return_value=[]), mock.patch(
            "knowledge.library._ocr_available",
            return_value=True,
        ), mock.patch(
            "knowledge.library._extract_pdf_blocks_via_ocr",
            side_effect=RuntimeError("已安装 OCR Python 依赖，但系统未安装 Tesseract 可执行程序。"),
        ):
            with self.assertRaisesRegex(ValueError, "Tesseract"):
                library.parse_uploaded_document("scan.pdf", b"fake-pdf")


class KnowledgeManagementTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.original_paths = {
            "DATA_DIR": library.DATA_DIR,
            "STORE_PATH": library.STORE_PATH,
            "DRAFTS_DIR": library.DRAFTS_DIR,
            "DOCS_DIR": library.DOCS_DIR,
            "_INITIALIZED": library._INITIALIZED,
        }
        library.DATA_DIR = self.root / "data"
        library.STORE_PATH = library.DATA_DIR / "knowledge_store.json"
        library.DRAFTS_DIR = library.DATA_DIR / "knowledge_drafts"
        library.DOCS_DIR = self.root / "docs"
        library._INITIALIZED = False
        library.DATA_DIR.mkdir(parents=True, exist_ok=True)
        library.DOCS_DIR.mkdir(parents=True, exist_ok=True)
        store = {
            "knowledge_bases": [
                {
                    "id": library.BUILTIN_KB_ID,
                    "name": library.BUILTIN_KB_NAME,
                    "kind": "builtin",
                    "created_at": "2026-01-01T00:00:00+00:00",
                    "updated_at": "2026-01-01T00:00:00+00:00",
                },
                {
                    "id": "kb_custom",
                    "name": "旧知识库",
                    "kind": "custom",
                    "created_at": "2026-01-01T00:00:00+00:00",
                    "updated_at": "2026-01-01T00:00:00+00:00",
                },
            ],
            "documents": [
                {
                    "id": "doc_custom",
                    "source_key": "upload:test",
                    "knowledge_base_id": "kb_custom",
                    "knowledge_base_name": "旧知识库",
                    "source_type": "upload",
                    "filename": "dp.docx",
                    "title": "旧标题",
                    "summary": "动态规划",
                    "text": "动态规划关注状态转移。",
                    "blocks": [{"text": "动态规划关注状态转移。", "location": "段落 1"}],
                    "chunks": [
                        {
                            "source": "dp.docx",
                            "knowledge_base_id": "kb_custom",
                            "knowledge_base_name": "旧知识库",
                            "text": "动态规划关注状态转移。",
                            "location": "段落 1",
                            "chunk_index": 0,
                        }
                    ],
                    "created_at": "2026-01-01T00:00:00+00:00",
                    "updated_at": "2026-01-01T00:00:00+00:00",
                }
            ],
        }
        library.STORE_PATH.write_text(json.dumps(store, ensure_ascii=False), encoding="utf-8")
        self.rebuild_patch = mock.patch("knowledge.library.rebuild_artifacts", return_value={"faiss_status": "skipped"})
        self.rebuild_patch.start()

    def tearDown(self) -> None:
        self.rebuild_patch.stop()
        library.DATA_DIR = self.original_paths["DATA_DIR"]
        library.STORE_PATH = self.original_paths["STORE_PATH"]
        library.DRAFTS_DIR = self.original_paths["DRAFTS_DIR"]
        library.DOCS_DIR = self.original_paths["DOCS_DIR"]
        library._INITIALIZED = self.original_paths["_INITIALIZED"]
        self.tmp.cleanup()

    def test_rename_knowledge_base_updates_documents_and_chunks(self) -> None:
        result = library.rename_knowledge_base("kb_custom", "新知识库")
        detail = library.get_document_detail("doc_custom")

        self.assertEqual(result["knowledge_base"]["name"], "新知识库")
        assert detail is not None
        self.assertEqual(detail["knowledge_base_name"], "新知识库")
        self.assertEqual(detail["chunks"][0]["text"], "动态规划关注状态转移。")

    def test_rename_builtin_knowledge_base_is_allowed(self) -> None:
        result = library.rename_knowledge_base(library.BUILTIN_KB_ID, "算法基础库")

        self.assertEqual(result["knowledge_base"]["name"], "算法基础库")
        self.assertEqual(result["knowledge_base"]["kind"], "builtin")

    def test_delete_document_removes_it_from_store(self) -> None:
        result = library.delete_document("doc_custom")

        self.assertTrue(result["deleted"])
        self.assertIsNone(library.get_document_detail("doc_custom"))

    def test_update_document_rebuilds_text_blocks_and_chunks(self) -> None:
        result = library.update_document(
            "doc_custom",
            title="新标题",
            text="第一段：前缀和用于快速求区间和。\n\n第二段：动态规划用于状态转移。",
        )
        detail = result["document"]

        self.assertEqual(detail["title"], "新标题")
        self.assertIn("前缀和", detail["text"])
        self.assertGreaterEqual(detail["chunk_count"], 1)
        self.assertIn("前缀和", detail["chunks"][0]["text"])

    def test_confirm_new_kb_draft_ignores_stale_builtin_selection(self) -> None:
        draft_id = "draft_new_kb"
        draft = {
            "draft_id": draft_id,
            "file_name": "new.docx",
            "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "summary": "新知识库文档",
            "extracted_text": "这是新知识库里的内容。",
            "blocks": [{"text": "这是新知识库里的内容。", "location": "段落 1"}],
            "chunks": [
                {
                    "source": "new.docx",
                    "text": "这是新知识库里的内容。",
                    "location": "段落 1",
                    "chunk_index": 0,
                }
            ],
            "target": {
                "knowledge_base_id": "",
                "knowledge_base_name": "新建测试库",
            },
            "created_at": "2026-01-01T00:00:00+00:00",
            "confirmed_at": None,
        }
        library.DRAFTS_DIR.mkdir(parents=True, exist_ok=True)
        (library.DRAFTS_DIR / f"{draft_id}.json").write_text(json.dumps(draft, ensure_ascii=False), encoding="utf-8")

        result = library.confirm_upload_draft(
            draft_id=draft_id,
            knowledge_base_id=library.BUILTIN_KB_ID,
        )

        self.assertEqual(result["knowledge_base"]["name"], "新建测试库")
        self.assertNotEqual(result["knowledge_base"]["id"], library.BUILTIN_KB_ID)
        self.assertEqual(result["document"]["knowledge_base_name"], "新建测试库")

    def test_cancel_upload_draft_removes_pending_file(self) -> None:
        draft_id = "draft_cancel_me"
        draft = {
            "draft_id": draft_id,
            "file_name": "cancel.docx",
            "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "summary": "待取消草稿",
            "extracted_text": "这是一份待取消的草稿。",
            "blocks": [{"text": "这是一份待取消的草稿。", "location": "段落 1"}],
            "chunks": [],
            "target": {
                "knowledge_base_id": "kb_custom",
                "knowledge_base_name": "",
            },
            "created_at": "2026-01-01T00:00:00+00:00",
            "confirmed_at": None,
        }
        library.DRAFTS_DIR.mkdir(parents=True, exist_ok=True)
        draft_path = library.DRAFTS_DIR / f"{draft_id}.json"
        draft_path.write_text(json.dumps(draft, ensure_ascii=False), encoding="utf-8")

        result = library.cancel_upload_draft(draft_id=draft_id)

        self.assertTrue(result["cancelled"])
        self.assertFalse(draft_path.exists())

    def test_cancel_upload_draft_tombstones_when_file_delete_is_denied(self) -> None:
        draft_id = "draft_cancel_denied"
        draft = {
            "draft_id": draft_id,
            "file_name": "cancel.docx",
            "confirmed_at": None,
        }
        library.DRAFTS_DIR.mkdir(parents=True, exist_ok=True)
        draft_path = library.DRAFTS_DIR / f"{draft_id}.json"
        draft_path.write_text(json.dumps(draft, ensure_ascii=False), encoding="utf-8")

        with mock.patch("pathlib.Path.unlink", side_effect=PermissionError("denied")):
            result = library.cancel_upload_draft(draft_id=draft_id)

        self.assertTrue(result["cancelled"])
        self.assertTrue(draft_path.exists())
        tombstone = json.loads(draft_path.read_text(encoding="utf-8"))
        self.assertEqual(tombstone["draft_id"], draft_id)
        self.assertIn("deleted_at", tombstone)
        with self.assertRaises(FileNotFoundError):
            library._load_draft(draft_id)

    def test_builtin_knowledge_base_cannot_be_deleted(self) -> None:
        with self.assertRaises(ValueError):
            library.delete_knowledge_base(library.BUILTIN_KB_ID)


class RetrieverFallbackTests(unittest.TestCase):
    def setUp(self) -> None:
        self.original_store = retriever._STORE_CACHE
        self.signature_patch = mock.patch("agent.retriever._store_signature", return_value=None)
        self.signature_patch.start()
        retriever._STORE_CACHE = retriever.KnowledgeStore(
            chunks=[
                retriever.KnowledgeChunk(
                    knowledge_base_id="kb_test",
                    knowledge_base_name="测试知识库",
                    source="graph.md",
                    text="BFS 使用队列进行层序遍历。",
                    location="段落 1",
                ),
                retriever.KnowledgeChunk(
                    knowledge_base_id="kb_test",
                    knowledge_base_name="测试知识库",
                    source="dp.md",
                    text="动态规划关注状态转移。",
                    location="段落 2",
                ),
            ],
            model_name="test-model",
            faiss_index=None,
            signature=None,
        )

    def tearDown(self) -> None:
        retriever._STORE_CACHE = self.original_store
        self.signature_patch.stop()

    def test_retrieve_uses_lexical_fallback_without_faiss(self) -> None:
        results = retriever.retrieve_with_scores("BFS 队列", k=2)

        self.assertEqual(results[0].source, "graph.md")
        self.assertGreater(results[0].score, 0)

    def test_retrieve_lexical_fallback_matches_chinese_substrings(self) -> None:
        results = retriever.retrieve_with_scores("动态规划", k=2)

        self.assertEqual(results[0].source, "dp.md")
        self.assertGreaterEqual(results[0].score, 2)


class ChatFallbackTests(IsolatedSessionsMixin, unittest.TestCase):
    def test_chat_without_api_key_returns_local_answer_and_error_type(self) -> None:
        request = chain.ChatRequest(message="怎么讲前缀和？", history=[])

        with mock.patch.dict(os.environ, {"OPENAI_API_KEY": ""}, clear=False), mock.patch(
            "agent.chain.retrieve_with_scores",
            return_value=[],
        ):
            response = chain.chat(request)

        self.assertIn("当前还没有命中本地知识库", response.answer)
        self.assertEqual(response.error_type, "model_config")
        self.assertIn("OPENAI_API_KEY", response.error or "")
        self.assertIsNotNone(response.session_id)
        self.assertIsNotNone(sessions.get_session(str(response.session_id)))

    def test_low_confidence_retrieval_does_not_force_rag(self) -> None:
        request = chain.ChatRequest(message="怎么准备英语作文？", history=[])
        weak_chunk = retriever.RetrievedChunk(
            knowledge_base_id="kb_test",
            knowledge_base_name="测试知识库",
            source="graph.md",
            text="BFS 使用队列进行层序遍历。",
            location="段落 1",
            score=0.1,
            retrieval_mode="semantic",
        )

        with mock.patch.dict(os.environ, {"OPENAI_API_KEY": ""}, clear=False), mock.patch(
            "agent.chain.retrieve_with_scores",
            return_value=[weak_chunk],
        ):
            response = chain.chat(request)

        self.assertFalse(response.used_rag)
        self.assertEqual(response.sources, [])
        self.assertEqual(response.evidence, [])
        self.assertEqual(response.retrieval_mode, "semantic")
        self.assertGreater(response.rag_confidence, 0)
        self.assertIn("低于阈值", response.low_confidence_reason or "")


class UploadLimitTests(unittest.TestCase):
    def test_upload_limit_defaults_to_regular_document_size(self) -> None:
        with mock.patch.dict(os.environ, {"UPLOAD_MAX_BYTES": ""}, clear=False):
            self.assertEqual(app_module._upload_max_bytes(), 50 * 1024 * 1024)

    def test_oversized_upload_returns_413(self) -> None:
        with mock.patch.dict(os.environ, {"UPLOAD_MAX_BYTES": "10"}, clear=False):
            with self.assertRaises(HTTPException) as context:
                app_module._reject_oversized_upload(11)

        self.assertEqual(context.exception.status_code, 413)
        self.assertIn("文件过大", str(context.exception.detail))


if __name__ == "__main__":
    unittest.main()
